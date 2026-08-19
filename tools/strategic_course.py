from __future__ import annotations

import html
import json
import re
from pathlib import Path

from docx import Document

from build_course import (
    AUTHORITIES,
    FACTS,
    ROOT,
    SOURCE_INDEX,
    WORLDS,
    authority_link,
    clean_visible,
    extract_modules,
    source_links,
)


SOURCE_STRATEGY = Path("/Users/anshuman/Downloads/Indian_Contract_Law_Strategic_Simulation_Games.docx")

FIELD_LABELS = {
    "Doctrinal spine": "doctrine",
    "Strategic learning objective": "objective",
    "The system being modelled": "system",
    "Core loop, five to ten minutes": "core",
    "Lab layer": "lab",
    "Win and loss conditions": "win_loss",
    "Strategic payoff": "payoff",
    "Replay hook": "replay",
    "Build notes": "build",
    "Estimated play time": "time",
}

WORLD_ACTIONS = {
    1: [
        ("ENGINEER THE RECORD", "Buy or preserve the fact that strengthens the earliest legal link."),
        ("INVOKE THE STATUS", "Commit to the strongest present classification and force a response."),
        ("HOLD THE OPTION", "Delay commitment, conserve resources and keep an exit route alive."),
    ],
    2: [
        ("CLARIFY THE CHANNEL", "Strengthen proof of what was communicated, when and by whom."),
        ("SEND THE OPERATIVE ACT", "Use the live offer, acceptance or revocation rule now."),
        ("KEEP THE OFFER ALIVE", "Preserve bargaining room without making a destructive move."),
    ],
    3: [
        ("TRACE THE EXCHANGE", "Document the requested act, sequence and source of value."),
        ("STRUCTURE THE ROUTE", "Commit to the consideration rule or a precise statutory exception."),
        ("BANK THE CONCESSION", "Preserve commercial value while keeping legal options open."),
    ],
    4: [
        ("VERIFY CAPACITY", "Spend time proving capacity, age, need and the contracting moment."),
        ("COMMIT THE ASSET", "Use the strongest available capacity or restitution route now."),
        ("LIMIT EXPOSURE", "Conserve the estate or counterparty position before value moves."),
    ],
    5: [
        ("PRESERVE CONSENT EVIDENCE", "Build a record of pressure, knowledge, causation and choice."),
        ("ATTACK THE DEFECT", "Commit to the best free-consent classification and demand relief."),
        ("DELAY THE ELECTION", "Keep the transaction alive while preserving a later response."),
    ],
    6: [
        ("MAP THE LAWFUL ROUTE", "Separate consideration, object, restraint and severable promises."),
        ("TRIGGER THE INVALIDITY RULE", "Commit to the statutory prohibition or public-policy ground."),
        ("REDRAFT AROUND RISK", "Preserve the lawful objective while narrowing the dangerous term."),
    ],
    7: [
        ("BUILD THE TERM RECORD", "Preserve words, notice, reliance, interface and drafting context."),
        ("PRESS THE READING", "Commit to the classification or interpretation that serves the client."),
        ("REPAIR THE DRAFT", "Spend less leverage now to reduce incorporation or ambiguity risk."),
    ],
    8: [
        ("SEQUENCE PERFORMANCE", "Map readiness, order, tender, timing and the acts already completed."),
        ("PERFORM OR ELECT", "Take the operative contractual step and force the next obligation."),
        ("RESERVE RIGHTS", "Keep performing selectively while preserving compensation or exit."),
    ],
    9: [
        ("LOCK LOSS EVIDENCE", "Build the counterfactual, causation, notice and mitigation record."),
        ("ELECT THE REMEDY", "Commit to termination, performance, substitution or compensation."),
        ("MITIGATE NOW", "Sacrifice leverage to stop avoidable loss and control exposure."),
    ],
    10: [
        ("MAP STANDING", "Trace parties, beneficiaries, benefits and the source of the obligation."),
        ("CLAIM THE ROUTE", "Commit to the strongest contractual, statutory or restitutionary claim."),
        ("PRESERVE THE BENEFIT", "Protect value while avoiding an irreversible enforcement theory."),
    ],
}


def extract_strategy_specs() -> list[dict]:
    doc = Document(SOURCE_STRATEGY)
    paragraphs = [
        ((paragraph.style.name if paragraph.style else ""), clean_visible(paragraph.text))
        for paragraph in doc.paragraphs
    ]
    headings = [
        index
        for index, (style, text) in enumerate(paragraphs)
        if style == "Heading 2" and re.match(r"^\d{2}\s+", text)
    ]
    specs: list[dict] = []
    for heading_number, start in enumerate(headings):
        next_heading = headings[heading_number + 1] if heading_number + 1 < len(headings) else len(paragraphs)
        end = next_heading
        for index in range(start + 1, next_heading):
            if paragraphs[index][0] == "Heading 1":
                end = index
                break
        block = [text for _, text in paragraphs[start:end] if text and not text.startswith("Modules ")]
        match = re.match(r"^(\d{2})\s+(.+)$", block[0])
        if not match or len(block) < 3:
            continue
        spec = {
            "id": int(match.group(1)),
            "title": match.group(2),
            "game": block[1],
        }
        current: str | None = None
        for line in block[2:]:
            if line == "HOW IT PLAYS: THE LAB LAYER":
                continue
            matched = False
            for label, key in FIELD_LABELS.items():
                if line.startswith(label):
                    value = line[len(label):].strip(" :")
                    spec[key] = value
                    current = key
                    matched = True
                    break
            if not matched and current:
                spec[current] = f"{spec[current]} {line}".strip()
        missing = [value for value in FIELD_LABELS.values() if not spec.get(value)]
        if missing:
            raise RuntimeError(f"Strategic brief {spec['id']:02d} is missing: {missing}")
        specs.append(spec)
    if [spec["id"] for spec in specs] != list(range(1, 70)):
        raise RuntimeError("Strategic brief extraction did not produce Modules 01 through 69")
    return specs


def sentences(value: str) -> list[str]:
    return [item.strip() for item in re.split(r"(?<=[.!?])\s+", value) if len(item.strip()) > 35]


def compact(value: str, limit: int = 280) -> str:
    value = clean_visible(value)
    if len(value) <= limit:
        return value
    clipped = value[: limit - 1].rsplit(" ", 1)[0]
    return clipped + "..."


def round_prompts(spec: dict) -> list[str]:
    pool = sentences(spec["lab"])
    weights = {
        "must decide": 8, "whether": 6, "choose": 5, "decision": 4,
        "sharpest": 4, "first,": 3, "second,": 3, "third,": 3,
        "central tension": 3, "pressure": 2, "shock": 2, "round": 1,
        "failure mode": 1, "tempted": 1,
    }
    candidates = [
        (index, sentence)
        for index, sentence in enumerate(pool)
        if len(sentence) >= 70
        and not sentence.startswith(("is a ", "runs ", "sets a ", "puts the ", "gives the ", "places the "))
    ]
    ranked = sorted(
        candidates,
        key=lambda item: (-sum(weight for keyword, weight in weights.items() if keyword in item[1].lower()), item[0]),
    )
    selected_items: list[tuple[int, str]] = []
    for index, sentence in ranked:
        candidate = compact(sentence)
        if candidate not in [value for _, value in selected_items]:
            selected_items.append((index, candidate))
        if len(selected_items) == 4:
            break
    selected_items.sort(key=lambda item: item[0])
    selected = [value for _, value in selected_items]
    fallbacks = sentences(spec["system"] + " " + spec["core"])
    for sentence in fallbacks:
        if len(selected) == 4:
            break
        candidate = compact(sentence)
        if candidate not in selected:
            selected.append(candidate)
    return selected[:4]


def model_variables(spec: dict) -> list[str]:
    text = spec["system"] + " " + spec["build"]
    found: list[str] = []
    meter_match = re.search(r"meters?\s*\(([^)]+)\)", text, flags=re.I)
    if meter_match:
        found.extend(re.split(r",\s*|\s+and\s+", meter_match.group(1)))
    found.extend(re.sub(r"([a-z])([A-Z])", r"\1 \2", token).title() for token in re.findall(r"\b[a-z]+[A-Z][A-Za-z]+\b", text))
    found.extend(
        match.group(1)
        for match in re.finditer(
            r"\b([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+){0,2})\s+(?:index|meter|counter|budget|clock|score|ledger|timeline|status)",
            text,
        )
    )
    cleaned: list[str] = []
    for value in found:
        value = compact(value.strip(" .,:;()"), 30)
        if value and value.lower() not in {item.lower() for item in cleaned}:
            cleaned.append(value)
    fallbacks = ["Legal Position", "Proof", "Leverage", "Exposure"]
    for fallback in fallbacks:
        if len(cleaned) >= 4:
            break
        if fallback.lower() not in {item.lower() for item in cleaned}:
            cleaned.append(fallback)
    return cleaned[:4]


def opponent_policy(spec: dict) -> str:
    for sentence in sentences(spec["lab"]):
        lowered = sentence.lower()
        if any(word in lowered for word in ("claude", "opponent", "counterparty", "rival", "regulator", "market")):
            return compact(sentence, 320)
    return "The opposing side attacks the weakest proved element, then adapts to any repeated tactic."


def scenario_for(module_id: int, spec: dict) -> tuple[str, str, str]:
    if module_id in FACTS:
        return FACTS[module_id]
    return (
        "A promise made at a family lunch is currently only words, but the client wants it enforced.",
        "the earliest missing statutory link and the evidence available to fill it",
        "The promise is documented before performance, with a requested act and part-payment.",
    )


def module_html(meta: dict, spec: dict) -> str:
    module_id = spec["id"]
    scenario, key_fact, replay_fact = scenario_for(module_id, spec)
    authority = AUTHORITIES.get(module_id, "The statutory text is the primary authority; applications remain fact-sensitive.")
    links = source_links(meta["anchor"] + " " + spec["doctrine"])
    source_markup = "".join(
        f'<a href="{html.escape(url)}" target="_blank" rel="noopener noreferrer">{html.escape(name)}</a>'
        for name, url in links
    )
    if module_id in AUTHORITIES:
        source_markup += (
            f'<a href="{html.escape(authority_link(authority))}" target="_blank" '
            f'rel="noopener noreferrer">Locate the leading authority</a>'
        )
    previous_href = "index.html" if module_id == 1 else f"module-{module_id - 1:02d}.html"
    previous_label = "Course map" if module_id == 1 else f"Module {module_id - 1:02d}"
    next_href = "index.html" if module_id == 69 else f"module-{module_id + 1:02d}.html"
    next_label = "Return to course map" if module_id == 69 else f"Continue to Module {module_id + 1:02d}"
    actions = [
        {"label": label, "description": description}
        for label, description in WORLD_ACTIONS[meta["world"]]
    ]
    data = {
        "rounds": round_prompts(spec),
        "opponent": opponent_policy(spec),
        "actions": actions,
        "keyFact": clean_visible(key_fact),
        "replayFact": clean_visible(replay_fact),
        "winLoss": spec["win_loss"],
    }
    safe_data = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    citation_status = (
        "This brief contains an item marked UNVERIFIED. Confirm the authorised report and current treatment before relying on it."
        if "UNVERIFIED" in spec["doctrine"]
        else "The brief supplies a verified statutory route. Recheck authorised reports, later treatment and State amendments for live work."
    )
    title = html.escape(spec["title"])
    game = html.escape(spec["game"])
    doctrine = html.escape(spec["doctrine"])
    objective = html.escape(spec["objective"])
    system = html.escape(spec["system"])
    core = html.escape(spec["core"])
    win_loss = html.escape(spec["win_loss"])
    payoff = html.escape(spec["payoff"])
    replay = html.escape(spec["replay"])
    build = html.escape(spec["build"])
    duration = html.escape(spec["time"])
    scenario_html = html.escape(clean_visible(scenario))
    key_fact_html = html.escape(clean_visible(key_fact))
    authority_html = html.escape(clean_visible(authority))
    citation_status_html = html.escape(citation_status)
    model_strip = "".join(f'<span class="chip">MODEL VAR // {html.escape(value)}</span>' for value in model_variables(spec))
    return f'''<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="referrer" content="no-referrer">
<title>Module {module_id:02d} | {title}</title>
<style>
:root {{ color-scheme:light dark; --bg:light-dark(#f9f7ef,#0d1112); --fg:light-dark(#172024,#e9f7ee); --panel:light-dark(#fffdf5,#131b1c); --muted:light-dark(#5a6668,#9ab0aa); --line:light-dark(#172024,#78a78f); --green:light-dark(#08783e,#65e69b); --amber:light-dark(#a85c00,#ffc45e); --red:light-dark(#a52b28,#ff7770); --blue:light-dark(#075ea8,#78bdff); font-family:"Courier New",ui-monospace,monospace; }}
* {{ box-sizing:border-box; }} [hidden] {{ display:none!important; }} html {{ background:var(--bg); }}
body {{ margin:0; padding:10px; color:var(--fg); background:var(--bg); font-size:15px; line-height:1.5; }} button,a {{ font:inherit; }}
.cabinet {{ position:relative; width:min(1120px,100%); margin:auto; overflow:hidden; border:2px solid var(--line); border-radius:8px; background:var(--panel); box-shadow:7px 7px 0 color-mix(in srgb,var(--line) 30%,transparent); }}
.scanlines {{ position:absolute; inset:0; z-index:9; pointer-events:none; opacity:.04; background:repeating-linear-gradient(to bottom,transparent 0,transparent 3px,var(--fg) 4px); }}
.topbar {{ position:relative; z-index:1; display:flex; justify-content:space-between; gap:14px; padding:12px 14px; border-bottom:2px solid var(--line); background:color-mix(in srgb,var(--green) 10%,var(--panel)); }}
.brand,.kicker,.label {{ text-transform:uppercase; letter-spacing:.09em; }} .brand {{ font-weight:700; }} .muted,.small {{ color:var(--muted); }} .small {{ font-size:.82rem; }}
.progress {{ height:10px; border-bottom:1px solid var(--line); background:repeating-linear-gradient(to right,transparent 0,transparent calc(20% - 3px),var(--line) calc(20% - 3px),var(--line) 20%); }} .progress span {{ display:block; height:100%; width:20%; background:var(--green); transition:width .25s steps(5,end); }}
.hero {{ padding:20px 16px 17px; border-bottom:1px dashed var(--line); }} .kicker {{ color:var(--green); font-size:.78rem; }} h1,h2,h3,p {{ margin:0; }} h1 {{ max-width:900px; margin-top:5px; font-size:clamp(1.55rem,4vw,2.65rem); line-height:1.08; }} h2 {{ margin:5px 0 12px; font-size:clamp(1.3rem,3vw,2rem); line-height:1.15; }} h3 {{ margin-bottom:7px; font-size:1rem; }} p+p {{ margin-top:9px; }}
.chips,.controls,.sources {{ display:flex; flex-wrap:wrap; gap:8px; }} .chips {{ margin-top:12px; }} .chip {{ padding:4px 8px; border:1px solid var(--line); font-size:.78rem; }}
.screen {{ min-height:520px; }} .stage {{ display:none; padding:22px 16px; }} .stage.active {{ display:block; animation:enter .25s steps(4,end); }}
.grid {{ display:grid; grid-template-columns:minmax(0,1.2fr) minmax(260px,.8fr); gap:18px; align-items:start; }} .mode-grid {{ display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); gap:12px; margin-top:18px; }}
.panel,.mode,.docket,.terminal {{ padding:14px; border:1px solid var(--line); background:color-mix(in srgb,var(--green) 5%,var(--panel)); }} .mode {{ text-align:left; cursor:pointer; }} .mode:hover,.mode:focus-visible,.action:hover,.action:focus-visible,.btn:hover,.btn:focus-visible {{ outline:3px solid var(--amber); outline-offset:2px; }} .mode strong {{ display:block; margin-bottom:7px; color:var(--green); }}
.docket {{ margin:14px 0; border-left:5px solid var(--amber); }} .rule {{ border-left:5px solid var(--blue); }}
.choices {{ display:grid; gap:9px; margin:16px 0; }} .choice,.action {{ width:100%; padding:12px; color:var(--fg); text-align:left; border:1px solid var(--line); background:var(--panel); cursor:pointer; }} .choice strong,.action strong {{ display:block; margin-bottom:5px; color:var(--green); }} .choice.selected {{ color:var(--bg); background:var(--green); }}
.btn {{ display:inline-flex; align-items:center; justify-content:center; min-height:40px; padding:8px 13px; color:var(--fg); border:1px solid var(--line); background:var(--panel); text-decoration:none; cursor:pointer; }} .btn.primary {{ color:var(--bg); background:var(--fg); }} .controls {{ margin-top:16px; }}
.hud {{ display:grid; grid-template-columns:repeat(5,minmax(0,1fr)); gap:8px; margin:14px 0; }} .stat {{ padding:9px; border:1px solid var(--line); }} .stat b {{ display:block; margin-bottom:5px; font-size:.76rem; letter-spacing:.05em; }} .bar {{ height:8px; border:1px solid var(--line); }} .bar span {{ display:block; height:100%; background:var(--green); transition:width .25s steps(6,end); }} .stat.exposure .bar span {{ background:var(--red); }} .resource-value {{ font-size:1.35rem; color:var(--amber); }}
.round-head {{ display:flex; justify-content:space-between; gap:12px; align-items:start; }} .round-tag {{ padding:5px 8px; border:1px solid var(--amber); color:var(--amber); white-space:nowrap; }} .actions {{ display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:9px; margin:16px 0; }} .action:disabled {{ opacity:.42; cursor:not-allowed; }} .cost {{ display:block; margin-top:7px; color:var(--amber); font-size:.78rem; }}
.terminal {{ max-height:190px; overflow:auto; border-color:var(--green); font-size:.83rem; }} .terminal p {{ padding:4px 0; border-bottom:1px dotted color-mix(in srgb,var(--line) 35%,transparent); }} .terminal p::before {{ content:"> "; color:var(--green); }}
.feedback {{ min-height:58px; padding:12px; border:1px dashed var(--line); }} .feedback.good {{ border-color:var(--green); }} .feedback.bad {{ border-color:var(--red); }}
.law-grid {{ display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); gap:10px; margin:16px 0; }} .law-box {{ padding:13px; border:1px solid var(--line); }} .law-box strong {{ display:block; margin-bottom:6px; color:var(--green); }}
.result {{ display:inline-block; margin:16px 0 8px; padding:9px 13px; border:3px double var(--green); color:var(--green); font-size:1.35rem; }} .result.loss {{ border-color:var(--red); color:var(--red); }} .sources {{ margin-top:12px; }} .sources a {{ padding:6px 8px; color:var(--fg); border:1px solid var(--line); }}
.nav {{ display:flex; justify-content:space-between; gap:10px; padding:12px 14px; border-top:2px solid var(--line); }}
@keyframes enter {{ from {{ opacity:0; transform:translateX(8px); }} }}
@media(max-width:760px) {{ body {{ padding:5px; }} .grid,.mode-grid,.law-grid,.actions {{ grid-template-columns:1fr; }} .hud {{ grid-template-columns:repeat(2,minmax(0,1fr)); }} .resource {{ grid-column:1/-1; }} .topbar,.round-head,.nav {{ flex-wrap:wrap; }} .screen {{ min-height:620px; }} }}
@media(prefers-reduced-motion:reduce) {{ *,*::before,*::after {{ animation-duration:.01ms!important; transition-duration:.01ms!important; }} }}
</style>
</head>
<body>
<main class="cabinet" id="game-root">
  <div class="scanlines" aria-hidden="true"></div>
  <header class="topbar"><div><div class="brand">LAW MATTERS // STRATEGY LAB {module_id:02d}</div><div class="small">{html.escape(WORLDS[meta['world']])}</div></div><div class="small" aria-live="polite">SCREEN <strong id="screen-count">1/5</strong> | MODE <strong id="mode-label">SELECT</strong></div></header>
  <div class="progress" aria-label="Game progress"><span id="progress-fill"></span></div>
  <section class="hero"><div class="kicker">Module {module_id:02d} | {game}</div><h1>{title}</h1><p class="muted">{objective}</p><div class="chips"><span class="chip">CORE + LAB</span><span class="chip">ADAPTIVE OPPONENT</span><span class="chip">{duration}</span></div></section>
  <section class="screen">
    <section class="stage active" data-stage="0">
      <div class="grid"><div><div class="kicker">Boot menu</div><h2>Choose the depth of play</h2><p>The doctrine is the physics of this game. The Core teaches the control once. The Lab exposes its variables, timing, burdens and failure modes.</p><div class="mode-grid"><button class="mode" type="button" data-mode="core"><strong>CORE RUN</strong><span>{core}</span></button><button class="mode" type="button" data-mode="lab"><strong>STRATEGY LAB</strong><span>{system}</span></button></div></div><aside class="panel rule"><h3>Win condition</h3><p>{win_loss}</p><p class="small">The Lab is designed for replay. Repeating a move becomes less effective because the opposing side adapts.</p></aside></div>
    </section>
    <section class="stage" data-stage="1">
      <div class="kicker">Core run | one honest surprise</div><h2>Find the controlling lever</h2><div class="docket"><strong>CASE FILE</strong><p>{scenario_html}</p></div><div class="choices" role="group" aria-label="Core decision"><button class="choice" type="button" data-core="0"><strong>A // MAP THE DECISIVE FACT</strong>Preserve and test {key_fact_html}.</button><button class="choice" type="button" data-core="1"><strong>B // FOLLOW THE LABEL</strong>Treat the transaction's informal name as the legal result.</button><button class="choice" type="button" data-core="2"><strong>C // OPTIMISE ONLY FOR VALUE</strong>Ignore the statutory elements and take the most profitable immediate move.</button></div><div class="feedback" id="core-feedback">Choose the move you could defend from the statute and the record.</div><div class="controls"><button class="btn" type="button" id="core-lab" hidden>Continue into the Lab</button><button class="btn primary" type="button" id="core-debrief" hidden>Debrief Core run</button></div>
    </section>
    <section class="stage" data-stage="2">
      <div class="grid"><div><div class="kicker">Lab briefing | expose the parameters</div><h2>Enter {game}</h2><div class="docket"><strong>SYSTEM MODEL</strong><p>{system}</p></div><div class="chips">{model_strip}</div><p>{spec['lab']}</p><div class="controls"><button class="btn primary" type="button" id="start-lab">Load round 1</button></div></div><aside class="panel"><h3>Opponent policy</h3><p>{html.escape(opponent_policy(spec))}</p><h3 style="margin-top:14px">Resolver notes</h3><p class="small">{build}</p></aside></div>
    </section>
    <section class="stage" data-stage="3">
      <div class="round-head"><div><div class="kicker">Live simulation</div><h2 id="round-title">Round 1</h2></div><span class="round-tag" id="round-tag">SET POSITION</span></div>
      <div class="hud" aria-label="Simulation state"><div class="stat"><b>LEGAL POSITION</b><div class="bar"><span id="bar-position"></span></div><span id="value-position">46</span></div><div class="stat"><b>PROOF</b><div class="bar"><span id="bar-proof"></span></div><span id="value-proof">40</span></div><div class="stat"><b>LEVERAGE</b><div class="bar"><span id="bar-leverage"></span></div><span id="value-leverage">35</span></div><div class="stat exposure"><b>EXPOSURE</b><div class="bar"><span id="bar-exposure"></span></div><span id="value-exposure">28</span></div><div class="stat resource"><b>INTERVENTION POINTS</b><span class="resource-value" id="value-resource">10</span></div></div>
      <div class="docket"><strong>DECISION WINDOW</strong><p id="round-prompt"></p></div><div class="actions" id="action-list"><button class="action choice" type="button" data-action="0"></button><button class="action choice" type="button" data-action="1"></button><button class="action choice" type="button" data-action="2"></button></div><div class="feedback" id="round-feedback">Choose one move. Its cost, trade-off and opponent response resolve immediately.</div><div class="controls"><button class="btn primary" type="button" id="next-round" hidden>Advance simulation</button></div><div class="terminal" id="terminal" aria-live="polite"><p>Simulation loaded. The opponent is watching for repeated tactics.</p></div>
    </section>
    <section class="stage" data-stage="4">
      <div class="kicker">Law desk | inspectable debrief</div><h2>Why the system resolved this way</h2><div class="result" id="final-result">CORE RUN COMPLETE</div><p id="final-copy">The controlling legal fact was identified.</p><div class="law-grid"><div class="law-box"><strong>DOCTRINAL SPINE</strong><p>{doctrine}</p></div><div class="law-box"><strong>STRATEGIC PAYOFF</strong><p>{payoff}</p></div><div class="law-box"><strong>WIN AND LOSS LOGIC</strong><p>{win_loss}</p></div><div class="law-box"><strong>REPLAY HOOK</strong><p>{replay}</p></div><div class="law-box"><strong>AUTHORITY FOR STUDY</strong><p>{authority_html}</p></div><div class="law-box"><strong>CITATION STATUS</strong><p>{citation_status_html}</p></div></div><p class="small">Source check: statutory anchors link to official India Code or responsible ministry sources. Case references are learning signposts, not a substitute for checking the authorised report, subsequent treatment and any applicable State amendment.</p><div class="sources">{source_markup}</div><div class="controls"><button class="btn" type="button" id="restart">Replay this module</button><a class="btn primary" href="{next_href}">{html.escape(next_label)}</a></div>
    </section>
  </section>
  <nav class="nav" aria-label="Course navigation"><a class="btn" href="{previous_href}">{html.escape(previous_label)}</a><a class="btn" href="index.html">Course map</a></nav>
</main>
<script>
(() => {{
  "use strict";
  const DATA = {safe_data};
  const root = document.getElementById("game-root");
  const stages = Array.from(root.querySelectorAll("[data-stage]"));
  const state = {{ position:46, proof:40, leverage:35, exposure:28, resource:10 }};
  const costs = [3,2,1];
  const baseEffects = [
    {{ position:8, proof:18, leverage:-3, exposure:-6 }},
    {{ position:17, proof:3, leverage:13, exposure:12 }},
    {{ position:-2, proof:4, leverage:7, exposure:-10 }}
  ];
  const roundTags = ["SET POSITION","OPPONENT ADAPTS","FACT SHOCK","CLOSE THE FILE"];
  let stage = 0, round = 0, selectedCore = null, awaitingAdvance = false;
  const history = [];
  function clamp(value) {{ return Math.max(0, Math.min(100, Math.round(value))); }}
  function showStage(index) {{ stage=index; stages.forEach((node,i)=>node.classList.toggle("active",i===index)); document.getElementById("screen-count").textContent=`${{index+1}}/5`; document.getElementById("progress-fill").style.width=`${{(index+1)*20}}%`; if(index>0) root.querySelector(".screen").scrollIntoView({{behavior:"smooth",block:"start"}}); }}
  function setMode(mode) {{ document.getElementById("mode-label").textContent=mode.toUpperCase(); showStage(mode==="core"?1:2); }}
  function updateHud() {{ ["position","proof","leverage","exposure"].forEach(key=>{{ document.getElementById(`bar-${{key}}`).style.width=`${{state[key]}}%`; document.getElementById(`value-${{key}}`).textContent=String(state[key]); }}); document.getElementById("value-resource").textContent=String(state.resource); }}
  function log(message) {{ const line=document.createElement("p"); line.textContent=message; document.getElementById("terminal").prepend(line); }}
  function renderRound() {{ awaitingAdvance=false; document.getElementById("next-round").hidden=true; document.getElementById("round-title").textContent=`Round ${{round+1}} of 4`; document.getElementById("round-tag").textContent=roundTags[round]; document.getElementById("round-prompt").textContent=DATA.rounds[round]; document.getElementById("round-feedback").textContent="Choose one move. Its cost, trade-off and opponent response resolve immediately."; document.getElementById("round-feedback").className="feedback"; root.querySelectorAll("[data-action]").forEach((button,index)=>{{ const action=DATA.actions[index]; button.innerHTML=`<strong>${{String.fromCharCode(65+index)}} // ${{action.label}}</strong><span>${{action.description}}</span><span class="cost">COST ${{costs[index]}} IP</span>`; button.disabled=state.resource<costs[index]; }}); updateHud(); }}
  function resolveAction(index) {{ if(awaitingAdvance||state.resource<costs[index])return; state.resource-=costs[index]; const repeated=history.filter(value=>value===index).length; const factor=repeated?Math.max(.45,1-repeated*.25):1; const effect=baseEffects[index]; Object.keys(effect).forEach(key=>state[key]=clamp(state[key]+effect[key]*factor)); history.push(index); const candidates=["position","proof","leverage"]; const weakest=candidates.reduce((a,b)=>state[a]<=state[b]?a:b); state[weakest]=clamp(state[weakest]-6); let counter=`Opponent attacks your weakest vector: ${{weakest.toUpperCase()}} -6.`; if(state.exposure>55){{ state.position=clamp(state.position-4); state.exposure=clamp(state.exposure+5); counter+=" High exposure also damages legal position."; }} if(repeated) counter+=` Repetition reduced this move to ${{Math.round(factor*100)}}% effectiveness.`; const label=DATA.actions[index].label; document.getElementById("round-feedback").textContent=`${{label}} resolved. ${{counter}}`; document.getElementById("round-feedback").className="feedback good"; log(`R${{round+1}}: ${{label}}. ${{counter}}`); root.querySelectorAll("[data-action]").forEach(button=>button.disabled=true); updateHud(); awaitingAdvance=true; document.getElementById("next-round").hidden=false; document.getElementById("next-round").textContent=round===3?"Resolve win condition":"Advance simulation"; }}
  function finishLab() {{ const win=state.position>=50&&state.proof>=45&&state.exposure<=70&&new Set(history).size>=2; const result=document.getElementById("final-result"); result.textContent=win?"STRATEGIC WIN":"INSTRUCTIVE LOSS"; result.classList.toggle("loss",!win); document.getElementById("final-copy").textContent=win?`You ended with Position ${{state.position}}, Proof ${{state.proof}}, Leverage ${{state.leverage}} and Exposure ${{state.exposure}}. The line is defensible and did not depend on one repeated tactic.`:`Your final state was Position ${{state.position}}, Proof ${{state.proof}}, Leverage ${{state.leverage}} and Exposure ${{state.exposure}}. Compare the failure with the module's stated win and loss logic, then replay with a different sequence.`; showStage(4); }}
  root.addEventListener("click",event=>{{ const mode=event.target.closest("[data-mode]"); if(mode){{setMode(mode.dataset.mode);return;}} const core=event.target.closest("[data-core]"); if(core&&selectedCore===null){{ selectedCore=Number(core.dataset.core); root.querySelectorAll("[data-core]").forEach(node=>node.classList.toggle("selected",node===core)); const good=selectedCore===0; const feedback=document.getElementById("core-feedback"); feedback.textContent=good?`Correct. The operative route turns on ${{DATA.keyFact}}, not on the label or immediate commercial preference.`:"That move skips the controlling statutory fact. The legal system will resolve before the commercial preference does."; feedback.className=`feedback ${{good?"good":"bad"}}`; document.getElementById("core-lab").hidden=false; document.getElementById("core-debrief").hidden=false; return; }} const action=event.target.closest("[data-action]"); if(action){{resolveAction(Number(action.dataset.action));return;}} }});
  document.getElementById("core-lab").addEventListener("click",()=>showStage(2));
  document.getElementById("core-debrief").addEventListener("click",()=>showStage(4));
  document.getElementById("start-lab").addEventListener("click",()=>{{round=0;showStage(3);renderRound();}});
  document.getElementById("next-round").addEventListener("click",()=>{{if(!awaitingAdvance)return;if(round===3){{finishLab();return;}}round+=1;renderRound();}});
  document.getElementById("restart").addEventListener("click",()=>window.location.reload());
  updateHud();
}})();
</script>
</body>
</html>'''.replace("—", " - ")


def build_index(specs: list[dict]) -> None:
    source = SOURCE_INDEX if SOURCE_INDEX.exists() else ROOT / "index.html"
    text = source.read_text(encoding="utf-8")
    text = text.replace("contract-law-course-map.html", "index.html")
    text = text.replace('"contractos-promise-machine.html"', '"module-01.html"')
    text = text.replace("`modules/module-${padded(module.id)}.html`", "`module-${padded(module.id)}.html`")
    text = text.replace("modules/module-${padded(moduleId)}.html", "module-${padded(moduleId)}.html")
    text = text.replace("1 PLAYABLE · 68 PLANNED", "69 STRATEGIC GAMES")
    text = text.replace("69 PLAYABLE", "69 STRATEGIC GAMES")
    text = text.replace("proposed game mechanics", "Core and Lab simulation mechanics")
    text = text.replace("Module 01 is playable. The remaining page links are structured placeholders ready for later replacement.", "Every module includes a Core run and a deeper strategic Lab.")
    text = text.replace(
        "Each module assigns the learner a role, presents plausible choices, reveals the practical consequence before explaining the Indian law, and then changes one material fact so the learner must reason again.",
        "Each module offers a short Core run and a deeper strategic Lab with limited resources, visible legal state, adaptive opposition and instructive failure.",
    )
    text = text.replace(
        'if (module.id === 1) {\n        link.dataset.live = "true";\n      } else {\n        link.dataset.planned = "true";\n        link.dataset.moduleId = String(module.id);\n      }',
        'link.dataset.live = "true";',
    )
    text = text.replace(
        'game.textContent = `${module.game}${module.id === 1 ? " · PLAYABLE" : " · PAGE PLANNED"}`;',
        'game.textContent = `${module.game} · CORE + LAB`;',
    )
    for spec in specs:
        pattern = re.compile(rf'(\{{ id: {spec["id"]},[^\n]* game: ")[^"]*(" \}})')
        text = pattern.sub(lambda match: match.group(1) + spec["game"].replace('"', "'") + match.group(2), text)
    (ROOT / "index.html").write_text(text.replace("—", " - "), encoding="utf-8")


def build_strategic_course() -> None:
    meta_modules = extract_modules()
    specs = extract_strategy_specs()
    build_index(specs)
    manifest_modules = []
    for meta, spec in zip(meta_modules, specs, strict=True):
        target = ROOT / f"module-{spec['id']:02d}.html"
        target.write_text(module_html(meta, spec), encoding="utf-8")
        manifest_modules.append({
            **meta,
            "title": spec["title"],
            "game": spec["game"],
            "file": target.name,
            "authority": AUTHORITIES.get(spec["id"]),
            "sources": [name for name, _ in source_links(meta["anchor"] + " " + spec["doctrine"])],
            "core": spec["core"],
            "lab": spec["lab"],
            "win_loss": spec["win_loss"],
            "build_notes": spec["build"],
        })
    manifest = {
        "generated": "2026-08-03",
        "source": SOURCE_STRATEGY.name,
        "modules": manifest_modules,
    }
    (ROOT / "legal-source-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Built 69 strategic Core and Lab simulation pages plus the course map")


if __name__ == "__main__":
    build_strategic_course()
